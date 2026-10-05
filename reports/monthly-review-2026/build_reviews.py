"""Build evidence-backed August and September 2026 monthly reviews."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
PROJECT = "AI-Based Egg Tray Inventory Management System (EG-Counter)"


def add_text(doc, text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold = bold
    return p


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)
    return t


def photo_pair(doc, left, right, captions):
    t = doc.add_table(rows=2, cols=2)
    t.autofit = False
    for i, file in enumerate((left, right)):
        cell = t.cell(0, i)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ROOT / file), width=Inches(3.05))
        t.cell(1, i).text = captions[i]


def base(month):
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(.68)
    sec.left_margin = sec.right_margin = Inches(.70)
    styles = d.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = None
    header = sec.header.paragraphs[0]
    header.text = "M. Kumarasamy College of Engineering  |  Department of Computer Science and Business Systems"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.text = f"EG-Counter  •  {month} 2026  •  Evidence-based progress review"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)

    for _ in range(3):
        add_text(d, "")
    p = add_text(d, "MONTHLY PROGRESS REVIEW", "Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_text(d, f"{month.upper()} 2026", bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_text(d, "Project ID: MKCE/R&I/Consultancy/2026/001")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(d, "")
    p = add_text(d, PROJECT, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(d, "")
    p = add_text(d, "Prepared by", bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        "Rahul S (927623BCB041)",
        "Yasvanthpalani S (927623BCB063)",
        "Dharaneesh Kesavan (927623BCB007)",
        "Dharani S (927623BCB008)",
        "Mukund Balaji B (927623BCB033)",
    ):
        p = add_text(d, line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(d, "")
    p = add_text(d, "Guided by", bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ("Dr. V. Banupriya, HoD/CSBS", "Mr. T. Viswanath Kani, AP/CSBS"):
        p = add_text(d, line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(d, "")
    p = add_text(d, "Department of Computer Science and Business Systems\nM. Kumarasamy College of Engineering\nThalavapalayam, Karur\nReport prepared 1 October 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    return d


def august():
    d = base("August")
    d.add_heading("1. Monthly Progress Summary", 1)
    add_text(d, "A working Android and Cloudflare/Roboflow counting route was established and measured against ten visually reviewed warehouse photos. The evaluation exposed large view-dependent errors; the app withheld a final count for an inconsistent three-view scene.")
    add_text(d, "Current phase: model evaluation and mobile/backend integration. The 31 August measurements are a per-photo detection baseline, not certified egg-filled inventory.")
    d.add_heading("2. Work Completed This Month", 1)
    d.add_heading("2.1 Prototype and inference integration", 2)
    add_text(d, "The Android client sent photographs to a Cloudflare Worker, which called the deployed Roboflow RF-DETR Medium model projec-mutta/2. Direct Roboflow and Worker counts matched on all ten evaluated frames. The client parsed the backend response without changing the tray count.")
    d.add_heading("2.2 Ground-reference and evaluation work", 2)
    add_text(d, "Ten images received manual visible-layer references; four uncertain photos were excluded. The frozen baseline was 2/10 exact images (20.0%), 22.9 trays mean absolute error, and 58.43% mean per-image count agreement. Count agreement means 100 times (1 minus absolute count error divided by the reference); it is not detection precision or warehouse accuracy.")
    d.add_heading("2.3 Photo evidence and measured counts", 2)
    table(d, ["Saved photo / case", "Manual visual reference", "Model / mobile pathway", "Count agreement"], [
        ("img05, six-stack wall", "120", "124", "96.67%"),
        ("img04, brown two-stack scene", "32", "29", "90.62%"),
        ("img07, single evaluated frame", "21", "21", "100.00%"),
        ("All 10 accepted frames", "10 reviewed references", "2 exact; MAE 22.9", "58.43% mean"),
    ])
    add_text(d, "The numbers above are per-photo backend counts in the mobile app path. They are not accepted three-view inventory totals. The retained photos may have been captured before August; this table dates their evaluation, not capture.")
    photo_pair(d, "accuracy-evaluation/test-images/img05.jpeg", "accuracy-evaluation/detections/img05.jpg", ("Original evaluated photo: 120 manually visible tray layers.", "Saved V2 detection overlay: 124 per-photo detections."))
    d.add_heading("3. Challenges Faced", 1)
    for item in (
        "One 60-layer scene yielded 9, 12 and 93 across three views; the app withheld a final total and requested a rescan.",
        "Dense stacks caused duplicate boxes; distant or dim stacks caused misses.",
        "Visual layers do not certify concealed egg occupancy. The benchmark had no independent physical recount.",
    ):
        add_text(d, "• " + item)
    d.add_heading("4. Plan for September", 1)
    for item in (
        "Audit annotations and scene overlap across data splits.",
        "Compare angle, crop and threshold methods against per-stack references.",
        "Test quality prompts and multi-view counting on untouched, physically counted scenes.",
    ):
        add_text(d, "• " + item)
    d.add_heading("5. Overall Project Status", 1)
    table(d, ["Phase", "Status at 31 August", "Measured progress"], [
        ("Dataset and visual references", "Baseline established", "10 accepted / 4 excluded"),
        ("Model and backend", "Integrated; inaccurate on varied views", "2/10 exact; MAE 22.9"),
        ("Android integration", "Prototype operational", "Response parsing verified"),
        ("Exact warehouse inventory", "Not verified", "No physical acceptance set"),
    ])
    d.add_heading("5.1 Overall Progress", 2)
    add_text(d, "Prototype integration and baseline evaluation were complete. Overall project completion cannot be quantified. The next milestone is an independent physical scene recount with exact stack counts and rejection coverage.")
    add_text(d, "Evidence: accuracy-evaluation/report.md, summary.json, results.csv, ground-truth-exclusions.csv, and saved img05 original/overlay.")
    d.save(OUT / "August_2026_Monthly_Progress_Review.docx")


def september():
    d = base("September")
    d.add_heading("1. Monthly Progress Summary", 1)
    add_text(d, "September focused on five-stack scene analysis, annotation repair, perspective-corrected diagnostic crops and a separate layer-heatmap research route. The deployed Android/backend model remained projec-mutta/2. No validated automatic 3D inventory count or new production model was established.")
    add_text(d, "Current phase: controlled improvement and validation. A manually assisted total of 100 versus a 99-layer visual reference is close for one known photo, but per-stack errors and manual intervention prevent a 99% system-accuracy claim.")
    d.add_heading("2. Work Completed This Month", 1)
    d.add_heading("2.1 Five-stack visual audit and live backend check", 2)
    add_text(d, "The wide photo was manually reviewed as five stacks of 20, 20, 20, 20 and 19 visible layers (99 total). A side photo showed 19 layers, overlapping the same scene and therefore not additive. A fresh HTTP 200 production Worker call returned 76 for the wide photo and 19 for the side photo; the third diagnostic slot returned 87 for a second wide view. These were diagnostic per-photo counts, not an accepted calibrated triplet.")
    d.add_heading("2.2 Perspective correction and band experiment", 2)
    add_text(d, "Manually selected stack faces, rectified to flatter crops, returned 19/22/20/22/17 = 100. Only one of five stacks was exact; summed absolute stack error was seven (MAE 1.4). Independent band candidates were 12/11/19/19/10 and did not resolve the count. Automatic localization proposed eight regions for five stacks.")
    d.add_heading("2.3 Photo evidence and measured counts", 2)
    table(d, ["Saved photo / method", "Manual visual reference", "Model result", "Count agreement"], [
        ("Wide original, V2", "99", "76", "76.77%"),
        ("Side original, V2", "19", "19", "100.00% (one photo)"),
        ("Wide, 5 manual rectified crops", "99 total", "100 total", "98.99% total only"),
        ("Same five crops, stack-level", "20/20/20/20/19", "19/22/20/22/17", "1/5 exact stacks"),
    ])
    photo_pair(d, "reports/two-view-20260924/input-1.jpg", "reports/two-view-20260924/manual-corrections/numbered-1.jpg", ("Original five-stack photograph used for September diagnostics.", "Numbered visual review: 99 layer references, not a physical recount."))
    d.add_heading("2.4 Mobile scan result for a reported 90 tray scene", 2)
    add_text(d, "The supplied mobile screenshot shows a scene described by the user as six stacks of 15 trays, or 90 trays. The app displayed 89 left, 92 right and 92 straight per-photo detections using projec-mutta/2. Relative to the user-reported count, those are 98.89%, 97.78% and 97.78% count agreement respectively. The screenshot says COUNT NOT VERIFIED; it provides no accepted fused total. The 90-tray reference has not been independently checked against the original capture photos.")
    table(d, ["Reference", "Left", "Right", "Straight", "App outcome"], [
        ("90 user-reported", "89", "92", "92", "Count not verified"),
    ])
    d.add_page_break()
    d.add_heading("Figure 2 Mobile scan evidence", 2)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT / "mobile-scan-90-reference.png"), width=Inches(4.1))
    add_text(d, "Screenshot supplied by the user. The mobile application reports per-photo detections and withholds a verified inventory total.")
    d.add_page_break()
    d.add_heading("3. Challenges Faced", 1)
    for item in (
        "The near-correct total concealed offsetting stack errors; perspective correction still missed and duplicated individual tray layers.",
        "Bands may track eggs or tray rims and cannot be converted to trays by a fixed +1 rule.",
        "Two photos did not establish calibrated camera pose, physical stack identity or hidden egg occupancy; 3D reconstruction and warehouse total remained unverified.",
        "Training examples had duplicate/scene-leakage risk; 118 approximate corrections were assigned to TRAIN only, with no independent acceptance evidence.",
    ):
        add_text(d, "• " + item)
    d.add_heading("4. Plan for October", 1)
    for item in (
        "Collect unchanged-scene front, left and right captures with physical per-stack filled/empty recounts and camera-quality metadata.",
        "Train a controlled V5 candidate only after curated split and annotation review; compare with unchanged V2 on unseen scenes.",
        "Measure automatic stack localization, exact per-stack and scene rates, MAE, false acceptance and rescan coverage before APK/backend promotion.",
        "Keep lighting, blur, angle and cut-off warnings active; use bands/edges as supporting evidence rather than self-validating tray counts.",
    ):
        add_text(d, "• " + item)
    d.add_heading("5. Overall Project Status", 1)
    table(d, ["Phase", "Status at 30 September", "Measured progress"], [
        ("Manual visual annotations", "Known scene reviewed", "99 wide + 19 overlapping side"),
        ("Deployed V2 mobile/backend", "Unchanged; per-photo errors persist", "76/99 wide; 19/19 side"),
        ("Assisted crop analysis", "Diagnostic only", "100/99 total; 1/5 exact stacks"),
        ("Heatmap research", "Not production-ready", "3/11 exact best count-selected arm"),
        ("Exact warehouse inventory", "Not verified", "No physical holdout acceptance"),
    ])
    d.add_heading("5.1 Overall Progress", 2)
    add_text(d, "A 98.99% count-agreement figure applies only to the summed output of one manually cropped, known photo. It is not a measured 99% field accuracy rate. The September heatmap comparison was development-only (11 reviewed faces, two retained groups, zero of two groups exact); it was not deployed. Exact egg-filled inventory remains a future acceptance milestone.")
    add_text(d, "Evidence: reports/two-view-20260924/manual-corrections/README.md and fresh-backend.json; crop-experiment/README.md and results.json; reports/stack-heatmap-followup-20260928/README.md; PROGRESS.md; user-supplied mobile-scan-90-reference.png. Photo and screenshot capture dates are not asserted here.")
    d.save(OUT / "September_2026_Monthly_Progress_Review.docx")


if __name__ == "__main__":
    august()
    september()
